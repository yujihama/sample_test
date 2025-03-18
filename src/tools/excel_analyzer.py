"""
Excel解析ツール - Excelファイルの構造解析や特定データの抽出を行うツール
"""

import os
from typing import Dict, List, Any, Optional, Tuple, Union
import uuid
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from datetime import datetime

from loguru import logger
from src.core.tool_base import ToolBase, ToolResult
from src.core.config import settings


class ExcelAnalyzer(ToolBase):
    """Excel解析ツール"""
    
    def __init__(self, tool_id: str = None):
        """
        Excel解析ツールの初期化
        
        Args:
            tool_id: ツールID（省略時は自動生成）
        """
        super().__init__(tool_id)
        self.description = "Excel解析ツール - Excel/CSVファイルの構造解析や特定データの抽出を行う"
        self.version = "1.0.0"
        self.capabilities = [
            "analyze_structure",    # 構造解析
            "extract_data",         # データ抽出
            "compare_sheets",       # シート間データ照合
            "search_values"         # 値の検索
        ]
        
        # 一時ファイル保存先
        self.temp_dir = os.path.join(settings.UPLOAD_DIR, "excel_analysis")
        os.makedirs(self.temp_dir, exist_ok=True)
        
        logger.info(f"Excel解析ツール初期化完了: {self.tool_id}")
    
    def validate_params(self, params: Dict[str, Any]) -> bool:
        """
        パラメータのバリデーション
        
        Args:
            params: 検証するパラメータ
            
        Returns:
            bool: 検証結果
        """
        if "operation" not in params:
            logger.error("operation パラメータが必要です")
            return False
            
        operation = params["operation"]
        
        if operation not in self.capabilities:
            logger.error(f"サポートされていない操作です: {operation}")
            return False
            
        # ファイルパスの確認
        if "file_path" not in params:
            logger.error("file_path パラメータが必要です")
            return False
            
        # 操作別の必須パラメータ確認
        if operation == "extract_data":
            if "sheet_name" not in params:
                logger.error("sheet_name パラメータが必要です")
                return False
                
        elif operation == "compare_sheets":
            required = ["sheet1", "sheet2", "key_column"]
            if not all(k in params for k in required):
                logger.error(f"操作 {operation} には次のパラメータが必要です: {required}")
                return False
                
        elif operation == "search_values":
            if "search_value" not in params:
                logger.error("search_value パラメータが必要です")
                return False
        
        return True
    
    async def execute(self, params: Dict[str, Any]) -> ToolResult:
        """
        ツール実行のメイン処理
        
        Args:
            params: 実行パラメータ
            
        Returns:
            ToolResult: 実行結果
        """
        try:
            if not self.validate_params(params):
                return ToolResult(
                    tool_id=self.tool_id,
                    status="error",
                    data={},
                    error_message="パラメータが無効です"
                )
            
            operation = params["operation"]
            file_path = params["file_path"]
            
            # ファイルの存在確認
            if not os.path.exists(file_path):
                return ToolResult(
                    tool_id=self.tool_id,
                    status="error",
                    data={},
                    error_message=f"ファイルが見つかりません: {file_path}"
                )
            
            # 拡張子の確認
            if not file_path.endswith((".xlsx", ".xls", ".csv")):
                return ToolResult(
                    tool_id=self.tool_id,
                    status="error",
                    data={},
                    error_message="サポートされていないファイル形式です（.xlsx, .xls, .csvのみ）"
                )
            
            # 操作の実行
            if operation == "analyze_structure":
                result_data = await self._process_analyze_structure(file_path, params)
            elif operation == "extract_data":
                result_data = await self._process_extract_data(file_path, params)
            elif operation == "compare_sheets":
                result_data = await self._process_compare_sheets(file_path, params)
            elif operation == "search_values":
                result_data = await self._process_search_values(file_path, params)
            else:
                return ToolResult(
                    tool_id=self.tool_id,
                    status="error",
                    data={},
                    error_message=f"操作 {operation} は未実装です"
                )
            
            return ToolResult(
                tool_id=self.tool_id,
                status="success",
                data=result_data
            )
            
        except Exception as e:
            logger.error(f"Excel解析ツール実行エラー: {e}")
            return await self.handle_error(e, params)
    
    async def _process_analyze_structure(self, file_path: str, params: Dict[str, Any]) -> Dict[str, Any]:
        """Excelファイルの構造解析"""
        if file_path.endswith(".csv"):
            # CSVファイルの場合
            df = pd.read_csv(file_path, nrows=10)
            structure = {
                "file_type": "CSV",
                "columns": df.columns.tolist(),
                "row_count": len(pd.read_csv(file_path)),
                "column_count": len(df.columns),
                "sheets": []
            }
        else:
            # Excelファイルの場合
            wb = load_workbook(file_path, read_only=True)
            sheets = []
            
            for sheet_name in wb.sheetnames:
                df = pd.read_excel(file_path, sheet_name=sheet_name, nrows=10)
                sheets.append({
                    "name": sheet_name,
                    "columns": df.columns.tolist(),
                    "row_count": wb[sheet_name].max_row,
                    "column_count": wb[sheet_name].max_column
                })
            
            structure = {
                "file_type": "Excel",
                "sheets": sheets
            }
        
        return {
            "structure": structure,
            "file_path": file_path
        }
    
    async def _process_extract_data(self, file_path: str, params: Dict[str, Any]) -> Dict[str, Any]:
        """特定シートからのデータ抽出"""
        sheet_name = params["sheet_name"]
        start_row = params.get("start_row", 0)
        end_row = params.get("end_row")
        columns = params.get("columns")
        
        # CSVファイルの場合、sheet_nameは無視
        if file_path.endswith(".csv"):
            df = pd.read_csv(file_path)
        else:
            try:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
            except ValueError:
                return {
                    "error": True,
                    "message": f"シート '{sheet_name}' が見つかりません"
                }
        
        # 行の範囲指定
        if end_row is not None:
            df = df.iloc[start_row:end_row]
        else:
            df = df.iloc[start_row:]
            
        # 列の指定
        if columns:
            try:
                df = df[columns]
            except KeyError:
                return {
                    "error": True,
                    "message": f"指定された列の一部またはすべてが見つかりません: {columns}"
                }
        
        # 結果を辞書のリストに変換
        result_data = df.to_dict(orient="records")
        
        # 大きすぎるデータセットの場合は要約
        if len(result_data) > 100:
            summary = {
                "total_rows": len(result_data),
                "preview": result_data[:10],
                "message": "結果が大きすぎるため、最初の10行のみを表示しています"
            }
            return summary
        
        return {
            "extracted_data": result_data,
            "row_count": len(result_data),
            "file_path": file_path,
            "sheet_name": sheet_name
        }
    
    async def _process_compare_sheets(self, file_path: str, params: Dict[str, Any]) -> Dict[str, Any]:
        """シート間のデータ照合"""
        sheet1 = params["sheet1"]
        sheet2 = params["sheet2"]
        key_column = params["key_column"]
        
        # CSVファイルの場合はエラー
        if file_path.endswith(".csv"):
            return {
                "error": True,
                "message": "この操作はExcelファイル（.xlsx, .xls）のみサポートしています"
            }
        
        try:
            df1 = pd.read_excel(file_path, sheet_name=sheet1)
            df2 = pd.read_excel(file_path, sheet_name=sheet2)
        except ValueError as e:
            return {
                "error": True,
                "message": f"シートの読み込みに失敗しました: {e}"
            }
        
        # キー列の存在確認
        if key_column not in df1.columns or key_column not in df2.columns:
            return {
                "error": True,
                "message": f"キー列 '{key_column}' が両方のシートに存在しません"
            }
        
        # 各シートのキー値を取得
        keys1 = set(df1[key_column].astype(str))
        keys2 = set(df2[key_column].astype(str))
        
        # 比較結果
        only_in_sheet1 = keys1 - keys2
        only_in_sheet2 = keys2 - keys1
        common_keys = keys1.intersection(keys2)
        
        return {
            "comparison_results": {
                "sheet1": sheet1,
                "sheet2": sheet2,
                "key_column": key_column,
                "only_in_sheet1_count": len(only_in_sheet1),
                "only_in_sheet1": list(only_in_sheet1)[:10],  # 最初の10個のみ
                "only_in_sheet2_count": len(only_in_sheet2),
                "only_in_sheet2": list(only_in_sheet2)[:10],  # 最初の10個のみ
                "common_keys_count": len(common_keys),
                "match_percentage": round(len(common_keys) / max(len(keys1), len(keys2)) * 100, 2) if max(len(keys1), len(keys2)) > 0 else 0
            }
        }
    
    async def _process_search_values(self, file_path: str, params: Dict[str, Any]) -> Dict[str, Any]:
        """値の検索"""
        search_value = params["search_value"]
        sheet_name = params.get("sheet_name")
        case_sensitive = params.get("case_sensitive", False)
        exact_match = params.get("exact_match", False)
        
        results = []
        
        if file_path.endswith(".csv"):
            # CSVファイルの場合
            df = pd.read_csv(file_path)
            sheet_results = self._search_in_dataframe(df, search_value, "CSV", case_sensitive, exact_match)
            if sheet_results:
                results.extend(sheet_results)
        else:
            # Excelファイルの場合
            if sheet_name:
                # 特定のシートのみを検索
                try:
                    df = pd.read_excel(file_path, sheet_name=sheet_name)
                    sheet_results = self._search_in_dataframe(df, search_value, sheet_name, case_sensitive, exact_match)
                    if sheet_results:
                        results.extend(sheet_results)
                except ValueError:
                    return {
                        "error": True,
                        "message": f"シート '{sheet_name}' が見つかりません"
                    }
            else:
                # すべてのシートを検索
                xlsx = pd.ExcelFile(file_path)
                for sheet in xlsx.sheet_names:
                    df = pd.read_excel(file_path, sheet_name=sheet)
                    sheet_results = self._search_in_dataframe(df, search_value, sheet, case_sensitive, exact_match)
                    if sheet_results:
                        results.extend(sheet_results)
        
        return {
            "search_results": {
                "file_path": file_path,
                "search_value": search_value,
                "match_count": len(results),
                "matches": results
            }
        }
    
    def _search_in_dataframe(self, df: pd.DataFrame, search_value: str, sheet_name: str, 
                            case_sensitive: bool, exact_match: bool) -> List[Dict[str, Any]]:
        """データフレーム内で値を検索"""
        results = []
        
        # 各セルを確認
        for row_idx, row in df.iterrows():
            for col_name in df.columns:
                cell_value = row[col_name]
                
                # NaN値はスキップ
                if pd.isna(cell_value):
                    continue
                
                # 文字列に変換
                cell_str = str(cell_value)
                search_str = str(search_value)
                
                # 大文字小文字を区別しない場合
                if not case_sensitive:
                    cell_str = cell_str.lower()
                    search_str = search_str.lower()
                
                # 一致チェック
                if (exact_match and cell_str == search_str) or \
                   (not exact_match and search_str in cell_str):
                    results.append({
                        "sheet": sheet_name,
                        "row": int(row_idx) + 1,  # 1-indexedに変換
                        "column": col_name,
                        "value": str(cell_value)
                    })
        
        return results 