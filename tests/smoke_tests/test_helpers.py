"""
テスト用ヘルパーモジュール

テスト実行のためのユーティリティ関数を提供します。
"""

import os
import sys
import json
import uuid
import shutil
import traceback
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, List, Optional

# プロジェクトルートをパスに追加
root_dir = Path(__file__).parents[2].absolute()
sys.path.append(str(root_dir))

# 必要なインポート
from src.utils.db_manager import get_db
from src.models.repositories import SampleDataRepository, AuditProcedureRepository
from src.utils.logger import setup_logger

# ロガーの設定
logger = setup_logger("smoke_test_helpers")

# APIのベースURL - 実際には使用しないがコード互換性のために残す
API_BASE_URL = os.environ.get("API_BASE_URL", "http://localhost:8000")


def setup_test_environment():
    """
    テスト実行環境をセットアップ
    - テストディレクトリの作成
    - テスト用サンプルデータの作成
    - 必要なファイルの配置
    """
    print("疎通テスト環境をセットアップします...")
    
    # テスト用ディレクトリの作成
    os.makedirs(os.path.join(root_dir, "data", "test_samples"), exist_ok=True)
    os.makedirs(os.path.join(root_dir, "data", "output"), exist_ok=True)
    os.makedirs(os.path.join(root_dir, "uploads"), exist_ok=True)
    os.makedirs(os.path.join(root_dir, "logs", "agent"), exist_ok=True)
    
    # テスト終了時のクリーンアップ登録
    import atexit
    atexit.register(cleanup_test_environment)
    
    try:
        # データベース接続確認
        db = next(get_db())
        db.execute("SELECT 1")  # 接続テスト
        print("データベース接続成功")
    except Exception as e:
        print(f"データベース接続エラー: {e}")
        # テスト用にインメモリDBを使用するなどの対応
        # この部分はプロジェクトの設定に依存
    
    # テストデータのセットアップ
    setup_sample_data()
    setup_procedure_data()
    setup_tool_registry()
    
    print("テスト環境セットアップ完了")
    return True


def cleanup_test_environment():
    """テスト環境のクリーンアップ"""
    print("疎通テスト環境をクリーンアップしています...")
    
    # テスト用の出力ファイルを削除
    output_dir = os.path.join(root_dir, "data", "output")
    if os.path.exists(output_dir):
        for filename in os.listdir(output_dir):
            if filename.startswith("test_") or "_test_" in filename:
                try:
                    os.remove(os.path.join(output_dir, filename))
                except Exception as e:
                    print(f"ファイル削除エラー: {e}")
    
    print("テスト環境クリーンアップ完了")


def setup_sample_data():
    """テスト用のサンプルデータをセットアップ"""
    # テスト用のサンプルデータを作成
    sample_data = {
        "サンプルID-001": {
            "type": "expense_claim",
            "description": "出張精算書類",
            "files": [
                "出張申請書.pdf",
                "サンプルID-001_領収書.png",
                "精算承認画面.jpg"
            ],
            "metadata": {
                "申請者": "山田太郎",
                "申請日": "2025-02-15",
                "金額": 12500,
                "承認者": "鈴木部長"
            }
        },
        "サンプルID-005": {
            "type": "transaction_record",
            "description": "取引履歴データ",
            "files": [
                "サンプルID-005_取引履歴.xlsx"
            ],
            "metadata": {
                "取引ID": "T-20250214-0023",
                "申請金額": 820000,
                "承認金額": 780000,
                "承認者": "佐藤次郎",
                "承認日": "2025-02-15"
            }
        }
    }
    
    # データベースにサンプルデータを登録
    db = next(get_db())
    sample_repo = SampleDataRepository(db)
    
    for sample_id, data in sample_data.items():
        # サンプルデータがなければ追加
        existing = sample_repo.get_by_id(sample_id)
        if not existing:
            # ファイルパスを生成
            file_path = os.path.join("data", "test_samples", f"{sample_id}_main.dat")
            
            # 正しい形式でcreateメソッドを呼び出す
            sample_repo.create({
                "id": sample_id,
                "filename": f"{sample_id}.dat",
                "file_path": file_path,
                "file_size": 1024,
                "file_type": data["type"],
                "row_count": 10,
                "column_count": 5,
                "columns": json.dumps(["col1", "col2", "col3", "col4", "col5"]),
                "file_metadata": json.dumps(data["metadata"])
            })
            
    # テスト用画像ファイルの作成
    create_dummy_image_file(os.path.join(root_dir, "data", "test_samples", "サンプルID-001_領収書.png"))
    
    # テスト用Excelファイルの作成
    create_dummy_excel_file(os.path.join(root_dir, "data", "test_samples", "サンプルID-005_取引履歴.xlsx"))


def setup_procedure_data():
    """テスト用の監査手続きデータをセットアップ"""
    # テスト用の監査手続きデータ
    procedure_data = {
        "PROC-001": {
            "title": "出張精算の妥当性検証",
            "description": """
            出張申請・精算プロセスの妥当性を検証する。
            具体的には以下の点を確認する：
            1. 出張期間と領収書日付の整合性
            2. 申請金額と領収書金額の一致
            3. 承認者の適切性（権限者による承認）
            """,
            "required_evidence": [
                "出張申請書",
                "領収書（交通費、宿泊費など）",
                "精算承認画面"
            ],
            "verification_points": [
                "出張期間と領収書日付の整合性",
                "申請金額と領収書金額の一致",
                "承認者の適切性（権限者による承認）"
            ]
        },
        "PROC-005": {
            "title": "取引記録の検証",
            "description": """
            取引記録の整合性と承認プロセスの妥当性を検証する。
            具体的には以下の点を確認する：
            1. 取引履歴（申請・承認・実行）の整合性
            2. 取引金額の一貫性
            3. システムログとの時系列整合性
            """,
            "required_evidence": [
                "取引履歴データ",
                "システムログ",
                "承認記録"
            ],
            "verification_points": [
                "取引履歴（申請・承認・実行）の整合性",
                "取引金額の一貫性",
                "システムログとの時系列整合性"
            ]
        }
    }
    
    # データベースに監査手続きデータを登録
    db = next(get_db())
    procedure_repo = AuditProcedureRepository(db)
    
    for proc_id, data in procedure_data.items():
        # 監査手続きデータがなければ追加
        existing = procedure_repo.get_by_id(proc_id)
        if not existing:
            # required_evidence と verification_points をJSON形式に変換
            required_data_fields = json.dumps({
                "required_evidence": data["required_evidence"],
                "verification_points": data["verification_points"]
            })
            
            # 正しい形式でcreateメソッドを呼び出す
            procedure_repo.create({
                "id": proc_id,
                "title": data["title"],
                "description": data["description"],
                "required_data_fields": required_data_fields,
                "risk_areas": json.dumps(["financial", "compliance"])
            })


def create_dummy_image_file(file_path):
    """
    テスト用のダミー画像ファイルを作成
    
    Args:
        file_path: ファイルパス
    """
    try:
        # PILがインストールされていれば簡単なテスト画像を作成
        try:
            from PIL import Image, ImageDraw, ImageFont
            
            # 画像サイズ
            width, height = 800, 600
            
            # 画像を作成
            image = Image.new('RGB', (width, height), color=(255, 255, 255))
            draw = ImageDraw.Draw(image)
            
            # テキスト領域を描画
            draw.rectangle([(300, 500), (400, 600)], outline="black", fill="lightgray")
            
            # テキスト追加
            draw.text((100, 100), "サンプル領収書", fill="black")
            draw.text((100, 150), "金額: 12,500円", fill="black")
            draw.text((100, 200), "日付: 2025/02/15", fill="black")
            draw.text((100, 250), "宛先: 山田太郎", fill="black")
            draw.text((300, 350), "承認者: 鈴木部長", fill="black")
            
            # 署名エリア
            draw.text((320, 520), "承認", fill="black")
            
            # 保存
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            image.save(file_path)
            print(f"テスト画像ファイル作成: {file_path}")
            
        except ImportError:
            # PILがない場合はテキストファイルで代用
            with open(file_path, 'w') as f:
                f.write("This is a dummy image file for testing purposes.\n")
                f.write("Content: Sample receipt\n")
                f.write("Amount: 12,500 JPY\n")
                f.write("Date: 2025/02/15\n")
                f.write("To: Taro Yamada\n")
                f.write("Approved by: Bucho Suzuki\n")
            print(f"テキスト形式のダミー画像ファイル作成: {file_path}")
            
    except Exception as e:
        print(f"ダミー画像ファイル作成エラー: {e}")
        # 最低限のファイルを作成
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, 'w') as f:
            f.write("dummy image")


def create_dummy_excel_file(file_path):
    """
    テスト用のダミーExcelファイルを作成
    
    Args:
        file_path: ファイルパス
    """
    try:
        # openpyxlがインストールされていれば簡単なテストExcelを作成
        try:
            import openpyxl
            from openpyxl import Workbook
            
            # ワークブックを作成
            wb = Workbook()
            ws = wb.active
            ws.title = "取引履歴"
            
            # ヘッダー
            headers = ["取引ID", "イベント種別", "タイムスタンプ", "担当者", "申請金額", "承認金額", "ステータス", "コメント"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
            
            # データ行
            data = [
                ["T-20250214-0023", "申請", "2025-02-14T10:23:45", "鈴木一郎", 820000, "", "申請完了", ""],
                ["T-20250214-0023", "承認", "2025-02-15T09:12:30", "佐藤次郎", 820000, 780000, "条件付承認", "予算超過のため調整"],
                ["T-20250214-0023", "実行", "2025-02-16T14:05:22", "システム自動処理", "", 780000, "完了", ""]
            ]
            
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, cell_value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=cell_value)
            
            # 保存
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            wb.save(file_path)
            print(f"テストExcelファイル作成: {file_path}")
            
        except ImportError:
            # openpyxlがない場合はCSVで代用
            import csv
            
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            with open(file_path.replace('.xlsx', '.csv'), 'w', newline='') as f:
                writer = csv.writer(f)
                writer.writerow(["取引ID", "イベント種別", "タイムスタンプ", "担当者", "申請金額", "承認金額", "ステータス", "コメント"])
                writer.writerow(["T-20250214-0023", "申請", "2025-02-14T10:23:45", "鈴木一郎", "820000", "", "申請完了", ""])
                writer.writerow(["T-20250214-0023", "承認", "2025-02-15T09:12:30", "佐藤次郎", "820000", "780000", "条件付承認", "予算超過のため調整"])
                writer.writerow(["T-20250214-0023", "実行", "2025-02-16T14:05:22", "システム自動処理", "", "780000", "完了", ""])
            
            # 一応元のパスにもダミーファイルを作成
            with open(file_path, 'w') as f:
                f.write("This is a dummy Excel file for testing purposes.\n")
                f.write("Content: CSV version available.\n")
            print(f"CSVとテキスト形式のダミーExcelファイル作成: {file_path}")
            
    except Exception as e:
        print(f"ダミーExcelファイル作成エラー: {e}")
        # 最低限のファイルを作成
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, 'w') as f:
            f.write("dummy excel file")


def setup_tool_registry():
    """
    テスト用のツールレジストリをセットアップ
    """
    # 必要なインポート
    try:
        from src.tools.tool_registry import ToolRegistry
        print("ToolRegistryのインポート成功")
        
        # ツールクラスのインポート
        try:
            from src.tools.image_processor import ImageProcessor
            print("ImageProcessorのインポート成功")
        except ImportError as e:
            print(f"ImageProcessorのインポートエラー: {e}")
            # フォールバック
            from unittest.mock import MagicMock
            ImageProcessor = MagicMock()
        
        try:
            from src.tools.excel_processor import ExcelProcessor
            print("ExcelProcessorのインポート成功")
        except ImportError as e:
            print(f"ExcelProcessorのインポートエラー: {e}")
            # フォールバック
            from unittest.mock import MagicMock
            ExcelProcessor = MagicMock()
        
        try:
            from src.tools.data_validator import DataValidator
            print("DataValidatorのインポート成功")
        except ImportError as e:
            print(f"DataValidatorのインポートエラー: {e}")
            # フォールバック
            from unittest.mock import MagicMock
            DataValidator = MagicMock()
        
        # ツールレジストリのインスタンスを取得（シングルトン）
        tool_registry = ToolRegistry()
        
        # 既存のツールをクリア
        ToolRegistry._instance = None
        tool_registry = ToolRegistry()
        
        # テスト用のツールを登録
        try:
            tool_registry.register_tool("image_processor", ImageProcessor)
            print("image_processorツール登録成功")
        except Exception as e:
            print(f"image_processorツール登録エラー: {e}")
        
        try:
            tool_registry.register_tool("excel_processor", ExcelProcessor)
            print("excel_processorツール登録成功")
        except Exception as e:
            print(f"excel_processorツール登録エラー: {e}")
        
        try:
            tool_registry.register_tool("data_validator", DataValidator)
            print("data_validatorツール登録成功")
        except Exception as e:
            print(f"data_validatorツール登録エラー: {e}")
        
        # 登録されたツールを確認
        registered_tools = tool_registry.get_all_tools()
        print(f"登録済みツール: {', '.join(registered_tools)}")
        
        return tool_registry
        
    except Exception as e:
        print(f"ツールレジストリセットアップエラー: {e}")
        return None


def get_test_config():
    """
    テスト用の設定を返す
    
    Returns:
        Dict[str, Any]: テスト設定
    """
    return {
        "timeout": 120,  # タイムアウト（秒）- 実際のLLM実行に対応するため延長
        "polling_interval": 0.5,  # ポーリング間隔（秒）
        "max_retries": 3,  # 最大再試行回数
        "use_real_llm": True,  # 実際のLLMを使用するフラグ
        "llm_timeout": 60,  # LLM呼び出しのタイムアウト（秒）
    }


def save_agent_message_log(agent_id: str, message_data: Dict[str, Any], workflow_id: Optional[str] = None):
    """
    エージェントのメッセージをログファイルに直接保存する
    
    Args:
        agent_id: エージェントID
        message_data: メッセージデータ
        workflow_id: ワークフローID（省略可能）
    """
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        workflow_suffix = f"_{workflow_id}" if workflow_id else ""
        filename = f"logs/agent/message_{agent_id}{workflow_suffix}_{timestamp}.json"
        
        with open(filename, "w", encoding="utf-8") as f:
            json.dump(message_data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"エージェント {agent_id} のメッセージをファイル {filename} に保存しました")
        return filename
    except Exception as e:
        logger.error(f"エージェント {agent_id} のメッセージ保存中にエラーが発生: {str(e)}")
        return None


def save_workflow_state_log(workflow_id: str, workflow_state: Dict[str, Any]):
    """
    ワークフロー状態をログファイルに直接保存する
    
    Args:
        workflow_id: ワークフローID
        workflow_state: ワークフロー状態データ
    """
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"logs/agent/workflow_state_{workflow_id}_{timestamp}.json"
        
        with open(filename, "w", encoding="utf-8") as f:
            json.dump(workflow_state, f, ensure_ascii=False, indent=2)
        
        logger.info(f"ワークフロー {workflow_id} の状態をファイル {filename} に保存しました")
        return filename
    except Exception as e:
        logger.error(f"ワークフロー {workflow_id} の状態保存中にエラーが発生: {str(e)}")
        return None


def save_agent_logs(workflow_id: Optional[str] = None):
    """
    テスト実行時のエージェントとワークフロー情報をファイルシステムから直接保存する
    
    Args:
        workflow_id: 特定のワークフローID（省略可能）
    """
    # テスト実行時刻
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    try:
        # まずログ保存ディレクトリの存在を確認
        log_dir = os.path.join(root_dir, "logs", "agent")
        os.makedirs(log_dir, exist_ok=True)
        
        # デバッグ情報の保存
        debug_log_path = os.path.join(log_dir, f"debug_log_{timestamp}.txt")
        with open(debug_log_path, "w", encoding="utf-8") as f:
            f.write(f"ログ保存処理開始: {datetime.now().isoformat()}\n")
            f.write(f"ワークフローID: {workflow_id}\n")
            f.write(f"Python実行パス: {sys.executable}\n")
            f.write(f"プロジェクトルート: {root_dir}\n")
            f.write(f"ログディレクトリ: {log_dir}\n")
            f.write(f"カレントディレクトリ: {os.getcwd()}\n")
            
        try:
            # メッセージブローカーから直接データを取得
            from src.core.messaging import MessageBroker
            
            broker = MessageBroker()
            
            # ブローカーの状態を記録
            with open(debug_log_path, "a", encoding="utf-8") as f:
                f.write(f"\nメッセージブローカー情報:\n")
                f.write(f"メッセージキュー数: {len(broker.message_queue) if hasattr(broker, 'message_queue') else '不明'}\n")
                f.write(f"メッセージ履歴数: {len(broker.message_history) if hasattr(broker, 'message_history') else '不明'}\n")
            
            # メッセージ履歴を保存
            if hasattr(broker, 'message_history') and broker.message_history:
                messages_filename = os.path.join(log_dir, f"message_history_{timestamp}.json")
                
                # ワークフローIDでフィルタリング
                if workflow_id:
                    filtered_messages = [msg for msg in broker.message_history if msg.get('workflow_id') == workflow_id]
                else:
                    filtered_messages = broker.message_history
                
                # ファイルに保存
                with open(messages_filename, "w", encoding="utf-8") as f:
                    json.dump(filtered_messages, f, ensure_ascii=False, indent=2)
                
                logger.info(f"メッセージ履歴をファイル {messages_filename} に保存しました")
                
                # デバッグログに記録
                with open(debug_log_path, "a", encoding="utf-8") as f:
                    f.write(f"\nメッセージ履歴保存成功: {messages_filename}\n")
                    f.write(f"フィルタリング後のメッセージ数: {len(filtered_messages)}\n")
            else:
                # エラーや空のケースをログに記録
                with open(debug_log_path, "a", encoding="utf-8") as f:
                    f.write("\nメッセージ履歴が見つかりませんでした。\n")
                    if not hasattr(broker, 'message_history'):
                        f.write("ブローカーにmessage_historyプロパティがありません。\n")
                    elif not broker.message_history:
                        f.write("message_historyは空です。\n")
            
            # ワークフロー状態を保存（ワークフローIDが指定されている場合）
            if workflow_id:
                # ワークフロー状態ファイルの場所を確認
                workflow_state_path = os.path.join(root_dir, "data", "workflow_states", f"{workflow_id}.json")
                
                with open(debug_log_path, "a", encoding="utf-8") as f:
                    f.write(f"\nワークフロー状態ファイルパス: {workflow_state_path}\n")
                    f.write(f"ファイルの存在: {os.path.exists(workflow_state_path)}\n")
                
                if os.path.exists(workflow_state_path):
                    # 状態ファイルを読み込む
                    with open(workflow_state_path, "r", encoding="utf-8") as f:
                        workflow_state = json.load(f)
                    
                    # ログディレクトリに保存
                    save_workflow_state_log(workflow_id, workflow_state)
                    with open(debug_log_path, "a", encoding="utf-8") as f:
                        f.write(f"ワークフロー状態の保存に成功しました。\n")
                else:
                    # ワークフロー情報が見つからない場合は、簡易的な情報を保存
                    simple_state_path = os.path.join(log_dir, f"workflow_state_{workflow_id}_{timestamp}.json")
                    with open(simple_state_path, "w", encoding="utf-8") as f:
                        json.dump({
                            "workflow_id": workflow_id,
                            "created_at": timestamp,
                            "note": "ワークフロー状態ファイルが見つからなかったため、テスト用に生成"
                        }, f, ensure_ascii=False, indent=2)
                    
                    with open(debug_log_path, "a", encoding="utf-8") as f:
                        f.write(f"ワークフロー状態ファイルが見つからないため、簡易情報を保存: {simple_state_path}\n")
            
            # デバッグログを最終更新
            with open(debug_log_path, "a", encoding="utf-8") as f:
                f.write(f"\nログ保存処理完了: {datetime.now().isoformat()}\n")
            
            logger.info("エージェントのログ情報を保存しました")
        except Exception as e:
            # エラー詳細をデバッグログに記録
            with open(debug_log_path, "a", encoding="utf-8") as f:
                f.write(f"\nエラーが発生しました: {str(e)}\n")
                f.write(f"エラータイプ: {type(e).__name__}\n")
                f.write("スタックトレース:\n")
                f.write(traceback.format_exc())
            
            logger.error(f"エージェントログの保存中にエラーが発生: {str(e)}")
            
            # 最低限のエラー情報をファイルに保存
            error_log_path = os.path.join(log_dir, f"error_log_{timestamp}.txt")
            with open(error_log_path, "w", encoding="utf-8") as f:
                f.write(f"ログ保存中にエラー発生: {datetime.now().isoformat()}\n")
                f.write(f"エラー: {str(e)}\n")
                f.write(traceback.format_exc())
    except Exception as e:
        # 最終的なフォールバックエラー処理
        logger.error(f"ログ保存の全体処理でエラーが発生: {str(e)}")
        try:
            # 最低限のエラー情報を標準のログディレクトリに保存
            with open(f"logs/critical_error_{timestamp}.txt", "w", encoding="utf-8") as f:
                f.write(f"重大なエラーが発生: {datetime.now().isoformat()}\n")
                f.write(f"エラー: {str(e)}\n")
                f.write(traceback.format_exc())
        except:
            # もはや何もできない
            pass 